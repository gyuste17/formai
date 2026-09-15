import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenpyxlImage

LOGO_PATH = r"C:\Users\gyust\GY Antigravity\formai\public\logos\formAI\1-removebg-preview-trimmed.png"
OUTPUT_BASE = r"C:\Users\gyust\OneDrive - guillermoyuste.es\1. GYusteWebs\formai.es\4.1 Documentos Clientes FormAI"

# Color constants
COLOR_TEAL = RGBColor(41, 151, 170)    # #2997AA
COLOR_NAVY = RGBColor(15, 23, 42)     # #0F172A
COLOR_GRAY = RGBColor(100, 116, 139)  # #64748B
HEX_TEAL = "2997AA"
HEX_NAVY = "0F172A"
HEX_LIGHT_BG = "F8FAFC"
HEX_HEADER_BG = "0F172A"
HEX_BORDER = "CBD5E1"

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_with_logo(doc, title, subtitle):
    # Top header table with 2 columns: Logo on left, document title on right
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(4.5)
    
    cell_left = table.cell(0, 0)
    cell_right = table.cell(0, 1)
    
    # Remove borders from header table
    for cell in [cell_left, cell_right]:
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
        tcPr.append(tcBorders)
    
    # Left: Logo
    p_logo = cell_left.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists(LOGO_PATH):
        p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.8))
        
    # Right: Title & Subtitle
    p_title = cell_right.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_title = p_title.add_run(title.upper() + "\n")
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(13)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_NAVY
    
    run_sub = p_title.add_run(subtitle)
    run_sub.font.name = 'Arial'
    run_sub.font.size = Pt(9)
    run_sub.font.color.rgb = COLOR_TEAL
    
    # Add horizontal divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(6)
    p_div.paragraph_format.space_after = Pt(16)
    p_div_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="{HEX_TEAL}"/></w:pBdr>')
    p_div._p.get_or_add_pPr().append(p_div_border)

def style_section_heading(p, text):
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY

# -------------------------------------------------------------
# 1. FICHA TÉCNICA - PLANTILLA FORMAL FORMAI (DOCX)
# -------------------------------------------------------------
def create_ficha_tecnica_template():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    add_header_with_logo(
        doc,
        "Ficha Técnica de Acción Formativa",
        "Formación Programada para Empresas · Bonificación FUNDAE"
    )
    
    # Section 1: Datos de la empresa cliente
    style_section_heading(doc.add_paragraph(), "1. DATOS DE LA EMPRESA BENEFICIARIA (CLIENTE)")
    t_emp = doc.add_table(rows=3, cols=2)
    t_emp.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_emp.columns[0].width = Inches(3.2)
    t_emp.columns[1].width = Inches(3.5)
    
    fields_emp = [
        ("Razón Social:", "[Nombre de la empresa cliente]"),
        ("CIF / NIF:", "[B12345678]"),
        ("Persona de Contacto / RRHH:", "[Nombre y Cargo del Responsable]"),
        ("Email / Teléfono de Contacto:", "[email@empresa.com / 600 000 000]"),
        ("Dirección del Centro de Trabajo:", "[Calle, Nº, Código Postal, Ciudad]"),
        ("Cuenta Cotización Seg. Social (CCC):", "[28/123456789/00]")
    ]
    
    idx = 0
    for r in range(3):
        for c in range(2):
            cell = t_emp.cell(r, c)
            set_cell_background(cell, HEX_LIGHT_BG if r % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            label, val = fields_emp[idx]
            run_lbl = p.add_run(f"{label} ")
            run_lbl.font.bold = True
            run_lbl.font.size = Pt(9.5)
            run_lbl.font.name = 'Arial'
            run_val = p.add_run(val)
            run_val.font.size = Pt(9.5)
            run_val.font.name = 'Arial'
            run_val.font.color.rgb = COLOR_GRAY
            idx += 1
            
    # Section 2: Datos de la formación
    style_section_heading(doc.add_paragraph(), "2. PLANIFICACIÓN DE LA ACCIÓN FORMATIVA")
    t_curso = doc.add_table(rows=4, cols=2)
    t_curso.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_curso.columns[0].width = Inches(3.2)
    t_curso.columns[1].width = Inches(3.5)
    
    fields_curso = [
        ("Denominación del Curso:", "[Ej: Inteligencia Artificial y ChatGPT en el Trabajo]"),
        ("Nº de Horas y Modalidad:", "[Ej: 15 Horas · Aula Virtual / Presencial]"),
        ("Nº de Participantes Previstos:", "[Ej: 10 alumnos]"),
        ("Nº Acción / Grupo FUNDAE:", "[A determinar tras comunicación]"),
        ("Fechas de Impartición:", "[Ej: Del 10 al 25 de Octubre de 2026]"),
        ("Horario de las Sesiones:", "[Ej: Martes y Jueves de 10:00 a 12:30h]"),
        ("Lugar / Enlace Aula Virtual:", "[Instalaciones del cliente / Microsoft Teams]"),
        ("Nivel y Perfil Requerido:", "[Iniciación / Intermedio / Avanzado]")
    ]
    
    idx = 0
    for r in range(4):
        for c in range(2):
            cell = t_curso.cell(r, c)
            set_cell_background(cell, HEX_LIGHT_BG if r % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            label, val = fields_curso[idx]
            run_lbl = p.add_run(f"{label} ")
            run_lbl.font.bold = True
            run_lbl.font.size = Pt(9.5)
            run_lbl.font.name = 'Arial'
            run_val = p.add_run(val)
            run_val.font.size = Pt(9.5)
            run_val.font.name = 'Arial'
            run_val.font.color.rgb = COLOR_GRAY
            idx += 1

    # Section 3: Equipo Formador & Entidad Acreditada
    style_section_heading(doc.add_paragraph(), "3. EQUIPO DOCENTE Y MARCO ADMINISTRATIVO FUNDAE")
    t_doc = doc.add_table(rows=2, cols=2)
    t_doc.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_doc.columns[0].width = Inches(3.2)
    t_doc.columns[1].width = Inches(3.5)
    
    fields_doc = [
        ("Formador / Dirección Pedagógica:", "Guillermo Yuste Durán (FormAI)"),
        ("Email / Teléfono Formador:", "hola@formai.es · +34 609 269 480"),
        ("Entidad Organizadora Homologada:", "Full Equipe S.L. (Entidad Acreditada FUNDAE)"),
        ("Gestión y Notificación Oficial:", "Tramitación completa ante aplicativo SEPE/FUNDAE")
    ]
    
    idx = 0
    for r in range(2):
        for c in range(2):
            cell = t_doc.cell(r, c)
            set_cell_background(cell, "F0F9FF") # subtle blue
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            label, val = fields_doc[idx]
            run_lbl = p.add_run(f"{label} ")
            run_lbl.font.bold = True
            run_lbl.font.size = Pt(9.5)
            run_lbl.font.name = 'Arial'
            run_val = p.add_run(val)
            run_val.font.size = Pt(9.5)
            run_val.font.name = 'Arial'
            idx += 1

    # Section 4: Objetivos y Contenidos
    style_section_heading(doc.add_paragraph(), "4. OBJETIVOS FORMATIVOS")
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.left_indent = Inches(0.2)
    p_obj.paragraph_format.space_after = Pt(8)
    run_obj = p_obj.add_run("• Capacitar a los profesionales en el uso eficiente y seguro de las herramientas digitales.\n"
                            "• Aplicar automatizaciones y flujos de trabajo orientados a la productividad del día a día.\n"
                            "• Desarrollar casos prácticos y proyectos adaptados al sector y operativa del cliente.\n"
                            "• Adquirir criterios de análisis crítico, verificación de calidad y buenas prácticas.")
    run_obj.font.name = 'Arial'
    run_obj.font.size = Pt(9.5)

    style_section_heading(doc.add_paragraph(), "5. PROGRAMA FORMATIVO Y MÓDULOS")
    p_mod = doc.add_paragraph()
    p_mod.paragraph_format.left_indent = Inches(0.2)
    p_mod.paragraph_format.space_after = Pt(12)
    run_mod = p_mod.add_run("Módulo 1: Fundamentos, conceptos clave y configuración del entorno de trabajo.\n"
                            "Módulo 2: Técnicas esenciales, funciones avanzadas y metodología práctica.\n"
                            "Módulo 3: Automatización de tareas repetitivas y casos de uso del equipo.\n"
                            "Módulo 4: Proyecto aplicado a la empresa, resolución de dudas y plan de mejora continua.")
    run_mod.font.name = 'Arial'
    run_mod.font.size = Pt(9.5)
    
    # Footer disclaimer
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(20)
    run_foot = p_foot.add_run("FormAI · Coordinación pedagógica por Guillermo Yuste en colaboración con Full Equipe S.L. (Entidad organizadora acreditada ante FUNDAE).")
    run_foot.font.name = 'Arial'
    run_foot.font.size = Pt(8)
    run_foot.font.italic = True
    run_foot.font.color.rgb = COLOR_GRAY

    out_path = os.path.join(OUTPUT_BASE, "2. Datos curso y participantes", "Ficha Técnica y Planificación - Plantilla FormAI.docx")
    doc.save(out_path)
    print("Created:", out_path)

# -------------------------------------------------------------
# 2. RECIBÍ DE MATERIAL DIDÁCTICO (DOCX)
# -------------------------------------------------------------
def create_recibi_material_template():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    add_header_with_logo(
        doc,
        "Entrega y Recibí de Material Didáctico",
        "Justificante de Entrega · Formación Bonificada FUNDAE"
    )
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(14)
    run_intro = p_intro.add_run("En cumplimiento de lo dispuesto en la normativa reguladora del Sistema de Formación Profesional para el Empleo (Ley 30/2015 y Orden TAS/2307/2007), se expide el presente documento que acredita la entrega de los medios didácticos y recursos formativos a los participantes de la acción formativa.")
    run_intro.font.name = 'Arial'
    run_intro.font.size = Pt(9.5)
    run_intro.font.color.rgb = COLOR_NAVY

    # Table with Course & Participant Data
    t = doc.add_table(rows=5, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.columns[0].width = Inches(2.5)
    t.columns[1].width = Inches(4.2)
    
    fields = [
        ("Nombre de la Acción Formativa:", "[Nombre del Curso]"),
        ("Nº Acción / Nº Grupo FUNDAE:", "[Acción ___ / Grupo ___]"),
        ("Empresa Participante:", "[Razón Social de la Empresa]"),
        ("Nombre y Apellidos del Alumno/a:", "[Nombre y Apellidos del Participante]"),
        ("DNI / NIE del Alumno/a:", "[12345678X]")
    ]
    
    for r in range(5):
        cell_lbl = t.cell(r, 0)
        cell_val = t.cell(r, 1)
        set_cell_background(cell_lbl, HEX_LIGHT_BG)
        set_cell_margins(cell_lbl, top=120, bottom=120, left=120, right=120)
        set_cell_margins(cell_val, top=120, bottom=120, left=120, right=120)
        
        lbl_text, val_text = fields[r]
        p_l = cell_lbl.paragraphs[0]
        run_l = p_l.add_run(lbl_text)
        run_l.font.name = 'Arial'
        run_l.font.size = Pt(9.5)
        run_l.font.bold = True
        
        p_v = cell_val.paragraphs[0]
        run_v = p_v.add_run(val_text)
        run_v.font.name = 'Arial'
        run_v.font.size = Pt(9.5)

    style_section_heading(doc.add_paragraph(), "DECLARACIÓN DE RECEPCIÓN")
    p_dec = doc.add_paragraph()
    p_dec.paragraph_format.space_after = Pt(16)
    run_dec = p_dec.add_run("Declaro por la presente haber recibido de forma íntegra y gratuita, con anterioridad o al inicio del curso, el material didáctico oficial y los recursos formativos necesarios para el correcto seguimiento de la formación, incluyendo:")
    run_dec.font.name = 'Arial'
    run_dec.font.size = Pt(9.5)
    
    p_bullets = doc.add_paragraph()
    p_bullets.paragraph_format.left_indent = Inches(0.2)
    p_bullets.paragraph_format.space_after = Pt(24)
    run_b = p_bullets.add_run("✔ Manual / Guía didáctica digital del curso con contenidos teóricos y ejercicios paso a paso.\n"
                              "✔ Hojas de cálculo, plantillas y archivos de datos para la realización de las prácticas.\n"
                              "✔ Acceso a la plataforma y herramientas tecnológicas requeridas para la impartición.\n"
                              "✔ Material de consulta de referencia, atajos y enlaces a recursos complementarios.")
    run_b.font.name = 'Arial'
    run_b.font.size = Pt(9.5)

    # Signature Area Table
    t_sig = doc.add_table(rows=1, cols=2)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sig.columns[0].width = Inches(3.3)
    t_sig.columns[1].width = Inches(3.4)
    
    c1 = t_sig.cell(0, 0)
    c2 = t_sig.cell(0, 1)
    for c in [c1, c2]:
        set_cell_margins(c, top=140, bottom=140, left=120, right=120)
        tcPr = c._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/><w:left w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/><w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/></w:tcBorders>')
        tcPr.append(tcBorders)
    
    p1 = c1.paragraphs[0]
    p1.add_run("Fecha de Entrega: _____ / _____ / 2026\n\nFirma del/la Alumno/a:\n\n\n\n____________________________________").font.size = Pt(9)
    
    p2 = c2.paragraphs[0]
    p2.add_run("Formación coordinada por FormAI\nEntidad Organizadora Acreditada: Full Equipe S.L.\n\nFirma del Formador / Responsable:\n\n\n____________________________________").font.size = Pt(9)

    out_path = os.path.join(OUTPUT_BASE, "3. Asistencia y Material", "RECIBÍ DE MATERIAL - Plantilla FormAI.docx")
    doc.save(out_path)
    print("Created:", out_path)

# -------------------------------------------------------------
# 3. CONTROL DE ASISTENCIA / PARTE DE FIRMAS (DOCX)
# -------------------------------------------------------------
def create_control_asistencia_template():
    doc = docx.Document()
    for s in doc.sections:
        s.orientation = docx.enum.section.WD_ORIENT.LANDSCAPE
        s.page_width = Inches(11.69)
        s.page_height = Inches(8.27)
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)
        
    add_header_with_logo(
        doc,
        "Control de Asistencia y Parte Diario de Firmas",
        "Formación Programada para Empresas · Requisito Oficial FUNDAE"
    )
    
    # Metadata Table
    t_meta = doc.add_table(rows=2, cols=3)
    t_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_meta.columns[0].width = Inches(3.4)
    t_meta.columns[1].width = Inches(3.4)
    t_meta.columns[2].width = Inches(3.4)
    
    meta_info = [
        ("Curso:", "[Nombre de la Acción Formativa]"),
        ("Empresa Cliente:", "[Razón Social Cliente]"),
        ("Nº Acción / Grupo:", "[Acción ___ / Grupo ___]"),
        ("Fechas y Horario:", "[Días concretos y franja horaria]"),
        ("Formador / Docente:", "Guillermo Yuste Durán (FormAI)"),
        ("Entidad Organizadora:", "Full Equipe S.L. (Acreditada FUNDAE)")
    ]
    
    idx = 0
    for r in range(2):
        for c in range(3):
            cell = t_meta.cell(r, c)
            set_cell_background(cell, HEX_LIGHT_BG)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            lbl, val = meta_info[idx]
            run_l = p.add_run(f"{lbl} ")
            run_l.font.bold = True
            run_l.font.size = Pt(8.5)
            run_l.font.name = 'Arial'
            run_v = p.add_run(val)
            run_v.font.size = Pt(8.5)
            run_v.font.name = 'Arial'
            idx += 1

    # Attendance Grid Table (10 students x 5 sessions)
    doc.add_paragraph().paragraph_format.space_before = Pt(10)
    
    cols_headers = ["Nº", "Apellidos y Nombre del Participante", "NIF / DNI", "Sesión 1\nFecha: ___", "Sesión 2\nFecha: ___", "Sesión 3\nFecha: ___", "Sesión 4\nFecha: ___", "Sesión 5\nFecha: ___", "Total Horas"]
    col_widths = [Inches(0.4), Inches(2.6), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1), Inches(0.7)]
    
    t_att = doc.add_table(rows=11, cols=len(cols_headers))
    t_att.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Style Header Row
    for c_idx, h_text in enumerate(cols_headers):
        cell = t_att.cell(0, c_idx)
        cell.width = col_widths[c_idx]
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=80, bottom=80, left=60, right=60)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.font.bold = True
        run.font.size = Pt(8)
        run.font.name = 'Arial'
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Style Student Rows
    for r in range(1, 11):
        for c in range(len(cols_headers)):
            cell = t_att.cell(r, c)
            cell.width = col_widths[c]
            set_cell_margins(cell, top=100, bottom=100, left=60, right=60)
            if r % 2 == 0:
                set_cell_background(cell, HEX_LIGHT_BG)
            if c == 0:
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(str(r)).font.size = Pt(8)
                
    p_note = doc.add_paragraph()
    p_note.paragraph_format.space_before = Pt(12)
    p_note.add_run("Nota obligatoria FUNDAE: Cada participante debe firmar personalmente en la casilla correspondiente a cada sesión diaria impartida. El formador certifica la asistencia real de los alumnos.").font.size = Pt(8)
    
    out_path = os.path.join(OUTPUT_BASE, "3. Asistencia y Material", "CONTROL DE ASISTENCIA - Plantilla FormAI.docx")
    doc.save(out_path)
    print("Created:", out_path)

# -------------------------------------------------------------
# 4. EXCEL DE DATOS DE LA BONIFICACIÓN (XLSX)
# -------------------------------------------------------------
def create_excel_participantes_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos Bonificación FormAI"
    ws.views.sheetView[0].showGridLines = True
    
    # Styles
    font_main_title = Font(name="Segoe UI", size=15, bold=True, color="0F172A")
    font_sub_title = Font(name="Segoe UI", size=10, italic=True, color="2997AA")
    font_sec_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_tbl_header = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
    font_data = Font(name="Segoe UI", size=9.5)
    font_bold = Font(name="Segoe UI", size=9.5, bold=True)
    
    fill_navy = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_teal = PatternFill(start_color="2997AA", end_color="2997AA", fill_type="solid")
    fill_light = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color="CBD5E1"),
        right=Side(style='thin', color="CBD5E1"),
        top=Side(style='thin', color="CBD5E1"),
        bottom=Side(style='thin', color="CBD5E1")
    )
    
    # Title Block
    ws['B2'] = "DATOS NECESARIOS PARA LA BONIFICACIÓN ANTE FUNDAE"
    ws['B2'].font = font_main_title
    ws['B3'] = "FormAI (Guillermo Yuste) · Gestión tramitada con entidad organizadora acreditada: Full Equipe S.L."
    ws['B3'].font = font_sub_title
    
    # Section 1: Empresa
    ws.merge_cells('B5:E5')
    ws['B5'] = "1. DATOS DE LA EMPRESA BENEFICIARIA (CLIENTE)"
    ws['B5'].font = font_sec_header
    ws['B5'].fill = fill_teal
    ws['B5'].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    
    emp_fields = [
        ("Razón Social:", "", "CIF / NIF:", ""),
        ("Dirección Completa:", "", "Código Postal y Población:", ""),
        ("Teléfono de Contacto:", "", "Email de Contacto RRHH:", ""),
        ("Convenio Colectivo:", "", "Código CNAE (4 dígitos):", ""),
        ("¿Dispone de RLT / Comité de Empresa?:", "No / Sí", "Cuenta Cotización Principal (CCC):", "")
    ]
    
    curr_row = 6
    for f1, v1, f2, v2 in emp_fields:
        ws.cell(curr_row, 2, f1).font = font_bold
        ws.cell(curr_row, 2).fill = fill_light
        ws.cell(curr_row, 3, v1).font = font_data
        ws.cell(curr_row, 4, f2).font = font_bold
        ws.cell(curr_row, 4).fill = fill_light
        ws.cell(curr_row, 5, v2).font = font_data
        
        for c in range(2, 6):
            ws.cell(curr_row, c).border = border_thin
            ws.cell(curr_row, c).alignment = Alignment(vertical="center")
        curr_row += 1

    # Section 2: Participantes
    curr_row += 2
    ws.merge_cells(f'B{curr_row}:M{curr_row}')
    ws[f'B{curr_row}'] = "2. DATOS DE LOS PARTICIPANTES A BONIFICAR (TRABAJADORES POR CUENTA AJENA)"
    ws[f'B{curr_row}'].font = font_sec_header
    ws[f'B{curr_row}'].fill = fill_navy
    ws[f'B{curr_row}'].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    
    curr_row += 1
    headers_part = [
        "Nº", "Primer Apellido", "Segundo Apellido", "Nombre", "NIF / NIE", 
        "Teléfono", "Email de Trabajo", "Nº Seg. Social (NASS)", "Fecha Nacimiento", 
        "Categoría Profesional", "Grupo Cotización (1-11)", "Cuenta Cotiz. CCC", "Nivel de Estudios"
    ]
    
    for c_idx, h_text in enumerate(headers_part, start=2):
        cell = ws.cell(curr_row, c_idx, h_text)
        cell.font = font_tbl_header
        cell.fill = fill_teal
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
        
    # Add 12 blank rows
    for r_idx in range(1, 13):
        curr_row += 1
        num_cell = ws.cell(curr_row, 2, r_idx)
        num_cell.font = font_bold
        num_cell.alignment = Alignment(horizontal="center", vertical="center")
        num_cell.border = border_thin
        if r_idx % 2 == 0:
            num_cell.fill = fill_light
            
        for c_idx in range(3, len(headers_part) + 2):
            cell = ws.cell(curr_row, c_idx, "")
            cell.border = border_thin
            cell.font = font_data
            if r_idx % 2 == 0:
                cell.fill = fill_light
                
    # Column widths adjustment
    col_widths = {
        'A': 4, 'B': 6, 'C': 18, 'D': 18, 'E': 16, 'F': 13, 
        'G': 14, 'H': 24, 'I': 20, 'J': 16, 'K': 22, 'L': 18, 'M': 20
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Add FormAI logo if image exists
    if os.path.exists(LOGO_PATH):
        try:
            img = OpenpyxlImage(LOGO_PATH)
            img.width = 140
            img.height = 78
            ws.add_image(img, 'L1')
        except Exception as e:
            print("Logo image insert note:", e)
            
    out_path = os.path.join(OUTPUT_BASE, "2. Datos curso y participantes", "Datos necesarios para la bonificación - FormAI.xlsx")
    wb.save(out_path)
    print("Created:", out_path)

# -------------------------------------------------------------
# 5. GUÍA INFORMATIVA DE FIRMA DE ENCOMIENDA (DOCX)
# -------------------------------------------------------------
def create_guia_encomienda():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)
        
    add_header_with_logo(
        doc,
        "Instrucciones: Adhesión a la Encomienda",
        "Autorización de Consulta de Crédito y Gestión · FUNDAE"
    )
    
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run_p = p.add_run("Estimado/a cliente:\n\n"
                      "Para que podamos gestionar la bonificación del curso y que su coste sea deducido al 100% de los seguros sociales de tu empresa, la normativa de FUNDAE (Ley 30/2015) exige formalizar el documento de Adhesión a la Encomienda de Gestión.")
    run_p.font.name = 'Arial'
    run_p.font.size = Pt(10)

    style_section_heading(doc.add_paragraph(), "¿QUÉ ES LA ENCOMIENDA DE GESTIÓN?")
    p_faq1 = doc.add_paragraph()
    p_faq1.paragraph_format.space_after = Pt(10)
    run_faq1 = p_faq1.add_run("Es la autorización administrativa estándar mediante la cual tu empresa autoriza a nuestra entidad organizadora oficial acreditada, FULL EQUIPE S.L., a consultar el saldo de crédito anual disponible ante el aplicativo del SEPE/FUNDAE y a comunicar telemáticamente el inicio y finalización del curso.")
    run_faq1.font.name = 'Arial'
    run_faq1.font.size = Pt(9.5)

    style_section_heading(doc.add_paragraph(), "¿TIENE ALGÚN COSTE O COMPROMISO?")
    p_faq2 = doc.add_paragraph()
    p_faq2.paragraph_format.space_after = Pt(10)
    run_faq2 = p_faq2.add_run("No. La consulta del crédito y la firma de la encomienda son totalmente gratuitas y no obligan a realizar ninguna formación si posteriormente decides no contratarla. Su única validez es habilitar el canal oficial de gestión con la Administración.")
    run_faq2.font.name = 'Arial'
    run_faq2.font.size = Pt(9.5)

    style_section_heading(doc.add_paragraph(), "PASOS PARA CUMPLIMENTAR EL DOCUMENTO:")
    steps = [
        ("1. Rellenar datos del representante legal:", "Nombre, apellidos y NIF del administrador o apoderado de la empresa."),
        ("2. Rellenar datos de la empresa cliente:", "Razón social completa, CIF y domicilio fiscal."),
        ("3. Indicar Cuenta de Cotización (CCC):", "La cuenta de cotización principal a la Seguridad Social."),
        ("4. Firma:", "Firma digital del representante legal (certificado digital de empresa) o firma manual con sello de la compañía."),
        ("5. Envío:", "Remitir el documento firmado a hola@formai.es para proceder a la comprobación del crédito.")
    ]
    
    for title, desc in steps:
        p_step = doc.add_paragraph()
        p_step.paragraph_format.left_indent = Inches(0.2)
        p_step.paragraph_format.space_after = Pt(4)
        r1 = p_step.add_run(f"• {title} ")
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r2 = p_step.add_run(desc)
        r2.font.size = Pt(9.5)

    p_contact = doc.add_paragraph()
    p_contact.paragraph_format.space_before = Pt(20)
    run_c = p_contact.add_run("¿Dudas con la cumplimentación? Contacta directamente con nosotros en hola@formai.es o al teléfono +34 609 269 480.")
    run_c.font.bold = True
    run_c.font.color.rgb = COLOR_TEAL
    run_c.font.size = Pt(9.5)

    out_path = os.path.join(OUTPUT_BASE, "1. Encomienda de Gestión - Consulta Crédito", "Guía de Firma - Encomienda de Gestión FUNDAE.docx")
    doc.save(out_path)
    print("Created:", out_path)

if __name__ == "__main__":
    create_ficha_tecnica_template()
    create_recibi_material_template()
    create_control_asistencia_template()
    create_excel_participantes_template()
    create_guia_encomienda()
    print("ALL CLIENT DOCUMENTS GENERATED SUCCESSFULLY!")
