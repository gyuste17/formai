import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from docx2pdf import convert

LOGO_PATH = r"C:\Users\gyust\GY Antigravity\formai\public\logos\formAI\1-removebg-preview-trimmed.png"
TARGET_DIR = r"C:\Users\gyust\OneDrive - guillermoyuste.es\1. GYusteWebs\formai.es\4.1 Documentos Clientes FormAI\1. Encomienda de Gestión - Consulta Crédito"

COLOR_TEAL = RGBColor(41, 151, 170)
COLOR_NAVY = RGBColor(15, 23, 42)
COLOR_GRAY = RGBColor(100, 116, 139)
HEX_TEAL = "2997AA"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"

def generate_encomienda():
    doc = docx.Document()
    
    # 1-page fit margins
    for s in doc.sections:
        s.top_margin = Inches(0.55)
        s.bottom_margin = Inches(0.55)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Header Table (Logo + Titles)
    header_tbl = doc.add_table(rows=1, cols=2)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tbl.autofit = False
    header_tbl.columns[0].width = Inches(2.2)
    header_tbl.columns[1].width = Inches(4.7)

    c_logo = header_tbl.cell(0, 0)
    c_text = header_tbl.cell(0, 1)

    for c in [c_logo, c_text]:
        tcPr = c._tc.get_or_add_tcPr()
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
        tcPr.append(borders)

    p_logo = c_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists(LOGO_PATH):
        p_logo.add_run().add_picture(LOGO_PATH, width=Inches(1.85))

    p_title = c_text.paragraphs[0]
    p_title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = p_title.add_run("DOCUMENTO DE ADHESIÓN\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(12)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_NAVY

    r2 = p_title.add_run("Contrato de Encomienda de Organización de la Formación\n")
    r2.font.name = "Arial"
    r2.font.size = Pt(9)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_TEAL

    r3 = p_title.add_run("Ley 30/2015, de 9 de septiembre · Real Decreto 694/2017")
    r3.font.name = "Arial"
    r3.font.size = Pt(8)
    r3.font.color.rgb = COLOR_GRAY

    # Teal Divider Line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(4)
    p_div.paragraph_format.space_after = Pt(10)
    p_bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="{HEX_TEAL}"/></w:pBdr>')
    p_div._p.get_or_add_pPr().append(p_bdr)

    # Intro Paragraph
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(8)
    r_intro = p_intro.add_run(
        "Documento de adhesión al Contrato de Encomienda de Organización de la Formación suscrito entre empresas "
        "al amparo de la Ley 30/2015, de 9 de septiembre, por la que se regula el Sistema de Formación Profesional para el Empleo "
        "en el ámbito laboral, formalizado entre la entidad organizadora acreditada FULL EQUIPE, S.L. y la empresa participante:"
    )
    r_intro.font.name = "Arial"
    r_intro.font.size = Pt(8.5)
    r_intro.font.color.rgb = COLOR_NAVY

    # Table: Datos de la Empresa y Representante
    tbl_data = doc.add_table(rows=4, cols=2)
    tbl_data.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_data.columns[0].width = Inches(3.45)
    tbl_data.columns[1].width = Inches(3.45)

    fields = [
        ("Nombre y Apellidos del Representante Legal:", "D./Dña. __________________________________________________"),
        ("NIF / NIE del Representante:", "_________________________"),
        ("Razón Social de la Empresa Beneficiaria:", "__________________________________________________"),
        ("CIF de la Empresa:", "_________________________"),
        ("Domicilio Social (Calle, Nº, CP, Municipio):", "__________________________________________________"),
        ("Provincia:", "_________________________"),
        ("Código Cuenta de Cotización (CCC Seg. Social):", "_________________________"),
        ("Facultad que ostenta (Administrador / Apoderado):", "__________________________________________________")
    ]

    f_idx = 0
    for r in range(4):
        for c in range(2):
            cell = tbl_data.cell(r, c)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_LIGHT_BG}"/>')
            cell._tc.get_or_add_tcPr().append(shd)

            # Cell borders
            bdr = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/><w:left w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/><w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/></w:tcBorders>')
            cell._tc.get_or_add_tcPr().append(bdr)

            # Internal margins
            tcMar = OxmlElement('w:tcMar')
            for m, val in [('top', 60), ('bottom', 60), ('left', 100), ('right', 100)]:
                n = OxmlElement(f'w:{m}')
                n.set(qn('w:w'), str(val))
                n.set(qn('w:type'), 'dxa')
                tcMar.append(n)
            cell._tc.get_or_add_tcPr().append(tcMar)

            p = cell.paragraphs[0]
            lbl, val = fields[f_idx]
            run_lbl = p.add_run(lbl + "\n")
            run_lbl.font.name = "Arial"
            run_lbl.font.size = Pt(7.5)
            run_lbl.font.bold = True
            run_lbl.font.color.rgb = COLOR_NAVY

            run_val = p.add_run(val)
            run_val.font.name = "Arial"
            run_val.font.size = Pt(8)
            run_val.font.color.rgb = COLOR_GRAY
            f_idx += 1

    # Heading: DECLARA
    p_dec = doc.add_paragraph()
    p_dec.paragraph_format.space_before = Pt(10)
    p_dec.paragraph_format.space_after = Pt(4)
    r_dec = p_dec.add_run("DECLARA Y MANIFIESTA:")
    r_dec.font.name = "Arial"
    r_dec.font.size = Pt(8.5)
    r_dec.font.bold = True
    r_dec.font.color.rgb = COLOR_NAVY

    clauses = [
        ("PRIMERO.", "Que la citada empresa está interesada y formaliza por el presente acto su adhesión al Contrato suscrito con la entidad externa organizadora e impartidora acreditada FULL EQUIPE, S.L. (con CIF B-85124428), en coordinación técnica y pedagógica con FormAI (Guillermo Yuste), para la organización, gestión y comunicación de la formación programada en dicha empresa al amparo de la Ley 30/2015 y Real Decreto 694/2017."),
        ("SEGUNDO.", "Que conoce y acepta el contenido íntegro de las condiciones, derechos y obligaciones fijadas en la normativa reguladora del Sistema de Formación Profesional para el Empleo en el ámbito laboral y en el referido contrato de encomienda."),
        ("TERCERO.", "Que por el presente documento acepta expresamente las obligaciones derivadas de dicha encomienda, autorizando a la entidad organizadora a la consulta del saldo de crédito formativo anual disponible ante el aplicativo oficial del SEPE / FUNDAE y a la comunicación telemática de las acciones formativas, surtiendo efectos desde el momento de su firma.")
    ]

    for num, txt in clauses:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.left_indent = Inches(0.15)
        p_c.paragraph_format.space_after = Pt(3)
        r_n = p_c.add_run(num + " ")
        r_n.font.name = "Arial"
        r_n.font.size = Pt(7.8)
        r_n.font.bold = True
        r_t = p_c.add_run(txt)
        r_t.font.name = "Arial"
        r_t.font.size = Pt(7.8)
        r_t.font.color.rgb = COLOR_NAVY

    # Date
    p_date = doc.add_paragraph()
    p_date.paragraph_format.space_before = Pt(8)
    p_date.paragraph_format.space_after = Pt(6)
    r_d = p_date.add_run("En _____________________________, a ______ de ___________________ de 2026.")
    r_d.font.name = "Arial"
    r_d.font.size = Pt(8.5)
    r_d.font.bold = True

    # Signature Table (Empresa + Full Equipe)
    tbl_sig = doc.add_table(rows=1, cols=2)
    tbl_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_sig.columns[0].width = Inches(3.45)
    tbl_sig.columns[1].width = Inches(3.45)

    c_s1 = tbl_sig.cell(0, 0)
    c_s2 = tbl_sig.cell(0, 1)

    for cell in [c_s1, c_s2]:
        tcPr = cell._tc.get_or_add_tcPr()
        bdr = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/><w:left w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/><w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/><w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER}"/></w:tcBorders>')
        tcPr.append(bdr)
        
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', 100), ('bottom', 100), ('left', 120), ('right', 120)]:
            n = OxmlElement(f'w:{m}')
            n.set(qn('w:w'), str(val))
            n.set(qn('w:type'), 'dxa')
            tcMar.append(n)
        tcPr.append(tcMar)

    p_s1 = c_s1.paragraphs[0]
    p_s1.add_run("POR LA EMPRESA BENEFICIARIA (CLIENTE)\n").font.size = Pt(8)
    p_s1.runs[0].font.bold = True
    p_s1.add_run("(Representante Legal / Apoderado)\n\n\n\n\nFirma y Sello de la Empresa:\n_______________________________________").font.size = Pt(7.5)

    p_s2 = c_s2.paragraphs[0]
    p_s2.add_run("POR LA ENTIDAD ORGANIZADORA ACREDITADA\n").font.size = Pt(8)
    p_s2.runs[0].font.bold = True
    p_s2.add_run("FULL EQUIPE, S.L. (CIF B-85124428)\n\n\n\n\nFirma y Sello de la Entidad:\n_______________________________________").font.size = Pt(7.5)

    # Footnote
    p_foot = doc.add_paragraph()
    p_foot.paragraph_format.space_before = Pt(8)
    p_foot.paragraph_format.space_after = Pt(0)
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_foot.add_run("FormAI · Coordinación pedagógica por Guillermo Yuste en colaboración con Full Equipe, S.L. (Entidad Homologada ante FUNDAE)")
    r_foot.font.name = "Arial"
    r_foot.font.size = Pt(7)
    r_foot.font.italic = True
    r_foot.font.color.rgb = COLOR_GRAY

    # Save docx
    out_docx = os.path.join(TARGET_DIR, "Encomienda de Gestión - FormAI - Full Equipe 2026.docx")
    doc.save(out_docx)
    print("Generated DOCX:", out_docx)

    # Convert to PDF
    try:
        convert(out_docx)
        print("Generated PDF:", out_docx.replace(".docx", ".pdf"))
    except Exception as e:
        print("PDF conversion note:", e)

if __name__ == "__main__":
    generate_encomienda()
